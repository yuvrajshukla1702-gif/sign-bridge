"""Convert docs/PROJECT_DOCUMENT.md → docs/PROJECT_DOCUMENT.docx"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "PROJECT_DOCUMENT.md"
OUT_PATH = ROOT / "PROJECT_DOCUMENT.docx"


def add_runs_with_bold(paragraph, content: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", content)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def flush_table(doc: Document, table_rows: list[str]) -> list[str]:
    if not table_rows:
        return []
    data: list[list[str]] = []
    for row in table_rows:
        cells = [c.strip() for c in row.strip("|").split("|")]
        if all(re.match(r"^:?-+:?$", c.replace(" ", "")) for c in cells):
            continue
        data.append(cells)
    if not data:
        return []
    cols = max(len(r) for r in data)
    table = doc.add_table(rows=len(data), cols=cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(data):
        for c_idx in range(cols):
            cell_text = row[c_idx] if c_idx < len(row) else ""
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs_with_bold(p, cell_text)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()
    return []


def convert() -> None:
    text = MD_PATH.read_text(encoding="utf-8")
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    lines = text.splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    table_rows: list[str] = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.strip().startswith("|") and "|" in line.strip()[1:]:
            table_rows.append(line)
            i += 1
            if i >= len(lines) or not lines[i].strip().startswith("|"):
                table_rows = flush_table(doc, table_rows)
            continue
        else:
            table_rows = flush_table(doc, table_rows)

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            i += 1
            continue

        if line.strip() in ("---", "***", "___"):
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        if line.strip().startswith(">"):
            content = line.strip()[1:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            add_runs_with_bold(p, content)
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        m_check = re.match(r"^[-*] \[ \] (.+)$", line)
        if m_check:
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_bold(p, "☐ " + m_check.group(1))
            i += 1
            continue

        m_bullet = re.match(r"^[-*] (.+)$", line)
        if m_bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_bold(p, m_bullet.group(1))
            i += 1
            continue

        m_num = re.match(r"^(\d+)\. (.+)$", line)
        if m_num:
            p = doc.add_paragraph(style="List Number")
            add_runs_with_bold(p, m_num.group(2))
            i += 1
            continue

        m_nested = re.match(r"^  [-*] (.+)$", line)
        if m_nested:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.5)
            add_runs_with_bold(p, m_nested.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs_with_bold(p, line.strip())
        i += 1

    flush_table(doc, table_rows)
    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")
    print(f"Size: {OUT_PATH.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    convert()
