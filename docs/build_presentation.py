"""
Generate an impressive project presentation (PPTX)
for the Sign Language Communication Bridge.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent / "PROJECT_PRESENTATION.pptx"

# Visual system — deep navy + teal (not purple / cream clichés)
NAVY = RGBColor(0x0B, 0x1C, 0x2C)
NAVY_2 = RGBColor(0x12, 0x2A, 0x40)
TEAL = RGBColor(0x1A, 0xA6, 0xA0)
TEAL_LT = RGBColor(0x5E, 0xD1, 0xCB)
CORAL = RGBColor(0xE8, 0x5D, 0x4C)  # emergency accent only
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xE8, 0xEE, 0xF2)
MUTED = RGBColor(0x9A, 0xB0, 0xC0)
CARD = RGBColor(0x16, 0x32, 0x4A)
GOLD = RGBColor(0xF0, 0xC4, 0x5C)

W, H = Inches(13.333), Inches(7.5)  # widescreen 16:9


def rgb_fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_bg(slide, color: RGBColor = NAVY) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    rgb_fill(shape, color)
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_accent_bar(slide, left, top, width, height=Inches(0.06), color=TEAL):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    rgb_fill(bar, color)
    return bar


def set_run(paragraph, text, size=18, bold=False, color=WHITE, font="Calibri"):
    paragraph.clear()
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return run


def add_text_box(slide, left, top, width, height, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p, text, size=size, bold=bold, color=color, font=font)
    return box


def card(slide, left, top, width, height, fill=CARD):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    rgb_fill(shape, fill)
    shape.adjustments[0] = 0.08
    return shape


def pill(slide, left, top, width, height, text, fill=TEAL, text_color=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    rgb_fill(shape, fill)
    shape.adjustments[0] = 0.5
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p, text, size=11, bold=True, color=text_color)
    shape.text_frame.paragraphs[0].space_before = Pt(0)
    return shape


def section_header(slide, title, subtitle=None):
    add_bg(slide)
    # left teal rail
    rail = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), H)
    rgb_fill(rail, TEAL)
    add_text_box(slide, Inches(0.55), Inches(0.28), Inches(11), Inches(0.5), title, size=28, bold=True, color=WHITE)
    add_accent_bar(slide, Inches(0.55), Inches(0.82), Inches(1.4))
    if subtitle:
        add_text_box(slide, Inches(0.55), Inches(0.95), Inches(12), Inches(0.4), subtitle, size=14, color=MUTED)


def footer(slide, page: int, total: int = 14):
    add_text_box(
        slide,
        Inches(0.55),
        Inches(7.05),
        Inches(10),
        Inches(0.3),
        "Sign Language Communication Bridge  ·  Team Project",
        size=10,
        color=MUTED,
    )
    add_text_box(
        slide,
        Inches(11.5),
        Inches(7.05),
        Inches(1.4),
        Inches(0.3),
        f"{page} / {total}",
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    total = 14

    # ─── 1. TITLE ───────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    add_bg(s, NAVY)
    # decorative right panel
    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.8), 0, Inches(4.533), H)
    rgb_fill(panel, NAVY_2)
    # teal diagonal-ish accents via stacked bars
    for i, y in enumerate([1.2, 1.45, 1.7, 1.95]):
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.8), Inches(y), Inches(0.18 + i * 0.35), Inches(0.12))
        rgb_fill(bar, TEAL if i % 2 == 0 else TEAL_LT)

    add_text_box(s, Inches(9.2), Inches(2.6), Inches(3.8), Inches(0.4), "REAL-TIME · TWO-WAY", size=12, bold=True, color=TEAL_LT)
    add_text_box(s, Inches(9.2), Inches(3.1), Inches(3.8), Inches(1.8), "Webcam signs become speech.\nVoice becomes sign visuals.\nEmergency help, one tap.", size=15, color=OFF_WHITE)

    pill(s, Inches(0.7), Inches(1.5), Inches(2.4), Inches(0.38), "ACCESSIBILITY TECH", fill=TEAL, text_color=NAVY)
    add_text_box(s, Inches(0.7), Inches(2.2), Inches(7.5), Inches(1.6), "Sign Language\nCommunication Bridge", size=40, bold=True, color=WHITE, font="Calibri")
    add_text_box(
        s,
        Inches(0.7),
        Inches(4.3),
        Inches(7.5),
        Inches(0.8),
        "Dynamic gesture recognition · LLM sentence polish · Speech ↔ Sign dialogue",
        size=16,
        color=MUTED,
    )
    add_accent_bar(s, Inches(0.7), Inches(5.2), Inches(2.2), Inches(0.07))
    add_text_box(s, Inches(0.7), Inches(5.5), Inches(7), Inches(0.4), "Team of 3  ·  Python · MediaPipe · PyTorch · Streamlit", size=14, color=OFF_WHITE)
    add_text_box(s, Inches(0.7), Inches(6.5), Inches(7), Inches(0.3), "Project Presentation  ·  2026", size=12, color=MUTED)

    # ─── 2. AGENDA ──────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "Agenda", "What we will cover today")
    items = [
        ("01", "The Problem", "Why static A–Z demos fall short"),
        ("02", "Our Solution", "Two-way bridge with temporal AI"),
        ("03", "Differentiators", "What makes this project stand out"),
        ("04", "Architecture & Stack", "How the system is built"),
        ("05", "Modules & Workflow", "From keypoints to spoken sentences"),
        ("06", "Research & Comparison", "How we chose the best approach"),
        ("07", "Team Plan & Roadmap", "3-person execution strategy"),
        ("08", "Demo & Impact", "What success looks like"),
    ]
    for i, (num, title, desc) in enumerate(items):
        col = i % 2
        row = i // 2
        left = Inches(0.55 + col * 6.3)
        top = Inches(1.55 + row * 1.25)
        card(s, left, top, Inches(6.0), Inches(1.05))
        add_text_box(s, left + Inches(0.25), top + Inches(0.18), Inches(0.7), Inches(0.4), num, size=20, bold=True, color=TEAL)
        add_text_box(s, left + Inches(1.0), top + Inches(0.18), Inches(4.7), Inches(0.35), title, size=16, bold=True, color=WHITE)
        add_text_box(s, left + Inches(1.0), top + Inches(0.55), Inches(4.7), Inches(0.35), desc, size=12, color=MUTED)
    footer(s, 2, total)

    # ─── 3. PROBLEM ─────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "The Problem", "Communication gaps need more than finger-spelling")
    problems = [
        ("Static A–Z only", "Most demos recognize letters, not real words or phrases used in conversation."),
        ("One-way systems", "Sign→text exists in isolation; hearing users rarely get a clear reverse path."),
        ("Broken grammar", "Gloss streams like “I go hospital” are hard for listeners without reconstruction."),
        ("No urgency path", "Emergency intent needs instant, high-confidence UI + audio — not a slow pipeline."),
    ]
    for i, (t, d) in enumerate(problems):
        left = Inches(0.55 + (i % 4) * 3.15)
        top = Inches(1.7)
        card(s, left, top, Inches(3.0), Inches(4.2))
        num_box = s.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(1.05), top + Inches(0.45), Inches(0.7), Inches(0.7))
        rgb_fill(num_box, TEAL)
        tf = num_box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p, str(i + 1), size=18, bold=True, color=NAVY)
        add_text_box(s, left + Inches(0.2), top + Inches(1.5), Inches(2.6), Inches(0.8), t, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text_box(s, left + Inches(0.2), top + Inches(2.4), Inches(2.6), Inches(1.5), d, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    footer(s, 3, total)

    # ─── 4. SOLUTION ────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "Our Solution", "A real-time two-way communication bridge")
    add_text_box(
        s,
        Inches(0.55),
        Inches(1.5),
        Inches(12.2),
        Inches(0.6),
        "Translate continuous dynamic gestures into natural spoken language — and convert voice back into sign visuals.",
        size=16,
        color=OFF_WHITE,
    )

    # Two pipeline cards
    card(s, Inches(0.55), Inches(2.3), Inches(6.0), Inches(3.8))
    pill(s, Inches(0.85), Inches(2.55), Inches(2.2), Inches(0.35), "SIGN  →  SPEECH", fill=TEAL, text_color=NAVY)
    steps_a = [
        "Webcam capture",
        "MediaPipe keypoints",
        "30-frame LSTM window",
        "Debounced word tokens",
        "LLM grammar polish",
        "Text-to-Speech output",
    ]
    box = s.shapes.add_textbox(Inches(0.85), Inches(3.15), Inches(5.4), Inches(2.7))
    tf = box.text_frame
    tf.word_wrap = True
    for i, step in enumerate(steps_a):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = f"{i+1}.  {step}"
        run.font.size = Pt(15)
        run.font.color.rgb = OFF_WHITE
        run.font.name = "Calibri"
        p.space_after = Pt(6)

    card(s, Inches(6.8), Inches(2.3), Inches(6.0), Inches(3.8))
    pill(s, Inches(7.1), Inches(2.55), Inches(2.4), Inches(0.35), "SPEECH  →  SIGN", fill=GOLD, text_color=NAVY)
    steps_b = [
        "Microphone input",
        "SpeechRecognition / Whisper",
        "Keyword extraction",
        "Sign GIF / animation match",
        "Split-screen visual reply",
        "Seamless dialogue loop",
    ]
    box = s.shapes.add_textbox(Inches(7.1), Inches(3.15), Inches(5.4), Inches(2.7))
    tf = box.text_frame
    for i, step in enumerate(steps_b):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = f"{i+1}.  {step}"
        run.font.size = Pt(15)
        run.font.color.rgb = OFF_WHITE
        run.font.name = "Calibri"
        p.space_after = Pt(6)
    footer(s, 4, total)

    # ─── 5. DIFFERENTIATORS ─────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "High-Impact Differentiators", "Four capabilities most student projects skip")
    diffs = [
        ("Dynamic Temporal Gestures", "Recognize multi-frame words & phrases — Thank you, Need help, Water — not only A–Z."),
        ("Two-Way Communication Loop", "Full dialogue: signs become speech; speech becomes sign visuals on screen."),
        ("LLM Sentence Reconstruction", "Gloss tokens → natural English via Ollama / Groq (with offline rule fallback)."),
        ("Emergency Quick-Action Mode", "Hotkeys for Call 108 / Medical Help — visual flash + priority audio."),
    ]
    for i, (t, d) in enumerate(diffs):
        top = Inches(1.5 + i * 1.25)
        card(s, Inches(0.55), top, Inches(12.2), Inches(1.1))
        badge = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top + Inches(0.28), Inches(0.7), Inches(0.55))
        rgb_fill(badge, CORAL if i == 3 else TEAL)
        badge.adjustments[0] = 0.2
        tf = badge.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p, f"0{i+1}", size=14, bold=True, color=NAVY if i != 3 else WHITE)
        add_text_box(s, Inches(1.8), top + Inches(0.18), Inches(10.5), Inches(0.35), t, size=18, bold=True, color=WHITE)
        add_text_box(s, Inches(1.8), top + Inches(0.55), Inches(10.5), Inches(0.4), d, size=14, color=MUTED)
    footer(s, 5, total)

    # ─── 6. TECH STACK ──────────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "Technology Stack", "Practical, laptop-friendly, demo-ready")
    stack = [
        ("Language", "Python 3.10+", "Entire pipeline in one runtime"),
        ("Vision", "OpenCV + MediaPipe", "Hands + upper-body landmarks"),
        ("Sequence AI", "PyTorch LSTM", "30-frame dynamic classifier"),
        ("Dashboard", "Streamlit", "Live video, chat, controls"),
        ("Sign → Voice", "pyttsx3 / gTTS", "Offline-first speech output"),
        ("Voice → Text", "SpeechRecognition\n+ Whisper", "Mic transcription"),
        ("Grammar LLM", "Ollama / Groq", "Gloss → natural sentence"),
        ("Data", "Custom .npy clips", "Your signs, your accuracy"),
    ]
    for i, (layer, tech, note) in enumerate(stack):
        col = i % 4
        row = i // 4
        left = Inches(0.55 + col * 3.15)
        top = Inches(1.55 + row * 2.5)
        card(s, left, top, Inches(3.0), Inches(2.25))
        add_text_box(s, left + Inches(0.2), top + Inches(0.25), Inches(2.6), Inches(0.3), layer.upper(), size=11, bold=True, color=TEAL)
        add_text_box(s, left + Inches(0.2), top + Inches(0.65), Inches(2.6), Inches(0.7), tech, size=16, bold=True, color=WHITE)
        add_text_box(s, left + Inches(0.2), top + Inches(1.45), Inches(2.6), Inches(0.55), note, size=12, color=MUTED)
    footer(s, 6, total)

    # ─── 7. ARCHITECTURE ────────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "System Architecture", "End-to-end data flow")

    nodes = [
        (0.4, "Webcam"),
        (2.15, "Landmarks"),
        (3.95, "30-Frame\nBuffer"),
        (5.75, "LSTM +\nDebounce"),
        (7.55, "Sentence\nBuffer"),
        (9.35, "LLM\nPolish"),
        (11.15, "TTS"),
    ]
    y = Inches(2.0)
    for i, (x, label) in enumerate(nodes):
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), y, Inches(1.55), Inches(1.0))
        rgb_fill(shape, TEAL if i in (0, 6) else CARD)
        shape.adjustments[0] = 0.15
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p, label, size=11, bold=True, color=NAVY if i in (0, 6) else WHITE)
        if i < len(nodes) - 1:
            arrow = s.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(x + 1.55),
                y + Inches(0.35),
                Inches(0.55),
                Inches(0.3),
            )
            rgb_fill(arrow, TEAL_LT)

    # reverse path
    add_text_box(s, Inches(0.55), Inches(3.4), Inches(12), Inches(0.35), "Reverse path (Speech → Sign)", size=14, bold=True, color=GOLD)
    rev = [
        (0.4, "Mic"),
        (2.5, "STT / Whisper"),
        (5.0, "Keyword Match"),
        (7.7, "Sign GIF Viewer"),
        (10.4, "Dialogue UI"),
    ]
    y2 = Inches(3.9)
    for i, (x, label) in enumerate(rev):
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), y2, Inches(1.9), Inches(0.75))
        rgb_fill(shape, RGBColor(0x3A, 0x2E, 0x14) if i % 2 else RGBColor(0x2A, 0x24, 0x12))
        # use gold border feel via gold fill on first/last
        if i in (0, 4):
            rgb_fill(shape, GOLD)
            tc = NAVY
        else:
            rgb_fill(shape, CARD)
            tc = OFF_WHITE
        tf = shape.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p, label, size=12, bold=True, color=tc)
        if i < len(rev) - 1:
            arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 1.9), y2 + Inches(0.22), Inches(0.5), Inches(0.3))
            rgb_fill(arrow, GOLD)

    card(s, Inches(0.55), Inches(5.0), Inches(12.2), Inches(1.5))
    add_text_box(s, Inches(0.85), Inches(5.2), Inches(11.5), Inches(0.35), "Feature design (per frame)", size=14, bold=True, color=TEAL)
    add_text_box(
        s,
        Inches(0.85),
        Inches(5.6),
        Inches(11.5),
        Inches(0.7),
        "21 hand landmarks × 2 hands × (x,y,z) + selected pose joints ≈ 147 features  ·  Wrist-relative + scale-normalized for distance invariance  ·  Predict only if confidence > 0.85 for ≥ 10 frames",
        size=13,
        color=OFF_WHITE,
    )
    footer(s, 7, total)

    # ─── 8. MODULES ─────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "Core Modules", "Clean ownership boundaries for a 3-person team")
    mods = [
        ("A", "Landmark Extraction", "MediaPipe Holistic\nNormalize to wrist / pose\nDistance-invariant features"),
        ("B", "Sequence Prediction", "Circular 30-frame buffer\nLSTM classifier\nDebounce + cooldown"),
        ("C", "Sentence + NLP", "Token accumulation\nIdle / EOS flush\nLLM grammar engine"),
        ("D", "Reverse Translation", "Mic → transcript\nKeyword → GIF DB\nSplit-screen viewer"),
    ]
    for i, (letter, title, body) in enumerate(mods):
        left = Inches(0.55 + i * 3.15)
        card(s, left, Inches(1.55), Inches(3.0), Inches(4.5))
        circle = s.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(1.05), Inches(1.85), Inches(0.85), Inches(0.85))
        rgb_fill(circle, TEAL)
        tf = circle.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p, letter, size=22, bold=True, color=NAVY)
        add_text_box(s, left + Inches(0.2), Inches(2.95), Inches(2.6), Inches(0.6), title, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        # body lines
        box = s.shapes.add_textbox(left + Inches(0.25), Inches(3.65), Inches(2.5), Inches(2.0))
        tf = box.text_frame
        tf.word_wrap = True
        for j, line in enumerate(body.split("\n")):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = line
            run.font.size = Pt(12)
            run.font.color.rgb = MUTED
            run.font.name = "Calibri"
            p.space_after = Pt(6)
    footer(s, 8, total)

    # ─── 9. COMPARISON ──────────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "Why This Approach Wins", "Research-backed comparison for best results")

    # table-like cards header
    headers = ["Approach", "Dynamic words", "Two-way", "Grammar LLM", "Laptop demo"]
    rows = [
        ["Static CNN A–Z", "No", "Rare", "No", "Easy"],
        ["Keypoints + SVM", "Weak", "Rare", "No", "Easy"],
        ["Raw video 3D CNN", "Yes", "Possible", "Rare", "Heavy GPU"],
        ["Our: MediaPipe + LSTM", "Yes", "Yes", "Yes", "Best fit"],
    ]
    # header row
    for c, h in enumerate(headers):
        left = Inches(0.55 + c * 2.5)
        shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.55), Inches(2.4), Inches(0.55))
        rgb_fill(shape, TEAL)
        tf = shape.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p, h, size=12, bold=True, color=NAVY)

    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            left = Inches(0.55 + c * 2.5)
            top = Inches(2.15 + r * 0.7)
            shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(2.4), Inches(0.65))
            highlight = r == 3
            rgb_fill(shape, RGBColor(0x1A, 0x4A, 0x48) if highlight else CARD)
            tf = shape.text_frame
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            set_run(p, val, size=12, bold=highlight or c == 0, color=TEAL_LT if highlight else OFF_WHITE)

    add_text_box(
        s,
        Inches(0.55),
        Inches(5.2),
        Inches(12.2),
        Inches(1.2),
        "Decision: Keypoints + LSTM gives strong accuracy without training huge video models. "
        "Add LLM polish and speech→sign GIFs for product-level differentiation. "
        "Keep a tiny vocabulary (8–12 signs) first, then expand.",
        size=14,
        color=MUTED,
    )
    footer(s, 9, total)

    # ─── 10. RESEARCH ───────────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "What We Research Online", "Learn these topics to outperform copy-paste demos")
    topics = [
        ("Computer Vision", "MediaPipe Hands vs Holistic\nLandmark normalization\nReal-time FPS on CPU"),
        ("Sequence Models", "LSTM / GRU / 1D-CNN\nSliding windows\nDebounce & smoothing"),
        ("Sign Linguistics", "ASL vs ISL choice\nGloss vs spoken English\nContinuous signing"),
        ("Speech & LLM", "Whisper sizes vs latency\nOllama vs Groq\nSafe prompt design"),
        ("Datasets", "WLASL · MS-ASL\nINCLUDE (ISL)\nCustom collection protocol"),
        ("HCI / Ethics", "Emergency UX patterns\nCamera consent\nLatency budgets"),
    ]
    for i, (t, body) in enumerate(topics):
        col = i % 3
        row = i // 3
        left = Inches(0.55 + col * 4.2)
        top = Inches(1.55 + row * 2.5)
        card(s, left, top, Inches(4.0), Inches(2.25))
        add_text_box(s, left + Inches(0.25), top + Inches(0.25), Inches(3.5), Inches(0.4), t, size=16, bold=True, color=TEAL)
        box = s.shapes.add_textbox(left + Inches(0.25), top + Inches(0.8), Inches(3.5), Inches(1.2))
        tf = box.text_frame
        tf.word_wrap = True
        for j, line in enumerate(body.split("\n")):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = "•  " + line
            run.font.size = Pt(13)
            run.font.color.rgb = OFF_WHITE
            run.font.name = "Calibri"
            p.space_after = Pt(4)
    footer(s, 10, total)

    # ─── 11. TEAM ───────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "Team Division — 3 Members", "Parallel tracks, shared integration contracts")
    roles = [
        ("Member A", "Vision & Data", "Landmarks · normalization\ndata_collection.py\nDataset quality & lighting", TEAL),
        ("Member B", "ML & Training", "Buffer · LSTM model\ntrain_model.py\nMetrics & checkpoint", GOLD),
        ("Member C", "App · Audio · NLP", "Streamlit UI · TTS/STT\nLLM polish · GIFs\nEmergency mode", CORAL),
    ]
    for i, (who, role, body, color) in enumerate(roles):
        left = Inches(0.55 + i * 4.2)
        card(s, left, Inches(1.55), Inches(4.0), Inches(4.5))
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.55), Inches(4.0), Inches(0.12))
        rgb_fill(bar, color)
        add_text_box(s, left + Inches(0.3), Inches(1.9), Inches(3.4), Inches(0.35), who, size=12, bold=True, color=color)
        add_text_box(s, left + Inches(0.3), Inches(2.35), Inches(3.4), Inches(0.5), role, size=20, bold=True, color=WHITE)
        box = s.shapes.add_textbox(left + Inches(0.3), Inches(3.1), Inches(3.4), Inches(2.4))
        tf = box.text_frame
        tf.word_wrap = True
        for j, line in enumerate(body.split("\n")):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = "▸  " + line
            run.font.size = Pt(14)
            run.font.color.rgb = OFF_WHITE
            run.font.name = "Calibri"
            p.space_after = Pt(8)
    footer(s, 11, total)

    # ─── 12. ROADMAP ────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "6-Week Roadmap", "From scaffold to demo day")
    weeks = [
        ("W1", "Research &\nfreeze vocab"),
        ("W2", "Collect data\n+ model scaffold"),
        ("W3", "Train LSTM\n+ live tokens"),
        ("W4", "LLM polish\n+ harden accuracy"),
        ("W5", "Speech→Sign\n+ UI polish"),
        ("W6", "Demo script\n+ final report"),
    ]
    for i, (w, desc) in enumerate(weeks):
        left = Inches(0.5 + i * 2.1)
        # connector
        if i < 5:
            line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(1.5), Inches(2.55), Inches(1.0), Inches(0.06))
            rgb_fill(line, TEAL)
        circle = s.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.45), Inches(2.25), Inches(0.7), Inches(0.7))
        rgb_fill(circle, TEAL if i < 5 else GOLD)
        tf = circle.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p, w, size=12, bold=True, color=NAVY)
        card(s, left, Inches(3.3), Inches(1.95), Inches(2.2))
        add_text_box(s, left + Inches(0.12), Inches(3.55), Inches(1.7), Inches(1.6), desc, size=13, color=OFF_WHITE, align=PP_ALIGN.CENTER)

    add_text_box(
        s,
        Inches(0.55),
        Inches(5.8),
        Inches(12.2),
        Inches(0.7),
        "Integration locks: feature dim = 147 · window = 30 frames · labels.json shared · GIF names = vocab names",
        size=14,
        color=MUTED,
    )
    footer(s, 12, total)

    # ─── 13. DEMO / IMPACT ──────────────────────────────────
    s = prs.slides.add_slide(blank)
    section_header(s, "Demo Script & Success Metrics", "Make the presentation unforgettable in 3 minutes")

    card(s, Inches(0.55), Inches(1.5), Inches(6.0), Inches(4.6))
    add_text_box(s, Inches(0.85), Inches(1.75), Inches(5.4), Inches(0.4), "Live demo flow", size=18, bold=True, color=TEAL)
    demo = [
        "Show live landmark overlay on webcam",
        "Sign: need_help + water → polished sentence + voice",
        "Speak: “thank you” → matching GIF",
        "Press Emergency → flash + priority audio",
        "Flash confusion-matrix / accuracy slide",
    ]
    box = s.shapes.add_textbox(Inches(0.85), Inches(2.4), Inches(5.4), Inches(3.3))
    tf = box.text_frame
    for i, item in enumerate(demo):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = f"{i+1}.  {item}"
        run.font.size = Pt(15)
        run.font.color.rgb = OFF_WHITE
        run.font.name = "Calibri"
        p.space_after = Pt(12)

    metrics = [
        ("≥ 85%", "Val accuracy\non isolated clips"),
        ("< 500ms", "Gesture → token\nlatency target"),
        ("8–12", "Signs in v1\nvocabulary"),
        ("100%", "Emergency path\nworks offline"),
    ]
    for i, (big, small) in enumerate(metrics):
        left = Inches(6.85 + (i % 2) * 3.1)
        top = Inches(1.5 + (i // 2) * 2.35)
        card(s, left, top, Inches(2.95), Inches(2.15))
        add_text_box(s, left + Inches(0.15), top + Inches(0.35), Inches(2.65), Inches(0.6), big, size=28, bold=True, color=TEAL_LT, align=PP_ALIGN.CENTER)
        add_text_box(s, left + Inches(0.15), top + Inches(1.15), Inches(2.65), Inches(0.7), small, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    footer(s, 13, total)

    # ─── 14. CLOSING ────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    add_bg(s, NAVY)
    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
    rgb_fill(panel, TEAL)
    add_text_box(s, Inches(0.9), Inches(1.8), Inches(11.5), Inches(0.5), "Thank You", size=44, bold=True, color=WHITE)
    add_accent_bar(s, Inches(0.9), Inches(2.55), Inches(2.0), Inches(0.08))
    add_text_box(
        s,
        Inches(0.9),
        Inches(2.9),
        Inches(11.5),
        Inches(1.0),
        "Real-time two-way sign language communication —\nbuilt to be researched, compared, and delivered as a team.",
        size=18,
        color=MUTED,
    )

    # three takeaway pills
    takes = ["Dynamic signs > A–Z", "Two-way + LLM polish", "Clear 3-person ownership"]
    for i, t in enumerate(takes):
        pill(s, Inches(0.9 + i * 3.7), Inches(4.3), Inches(3.4), Inches(0.5), t, fill=CARD, text_color=TEAL_LT)
        # override - need visible text on dark pill
        shape = s.shapes[-1]
        rgb_fill(shape, NAVY_2)
        # redraw text
        tf = shape.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p, t, size=13, bold=True, color=TEAL_LT)

    add_text_box(s, Inches(0.9), Inches(5.4), Inches(11.5), Inches(0.4), "Questions?  ·  Full write-up in PROJECT_DOCUMENT.docx", size=16, color=OFF_WHITE)
    add_text_box(s, Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.35), "Docs · Code scaffold · Report outline · Research checklist — ready for your team", size=13, color=MUTED)

    prs.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Size: {OUT.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    build()
